"""make_samples.py - builds sample.pdf and sample.docx to try pdf2md without your own documents.
Both contain what real documents have: a repeated page header/footer, a table of contents, headings, a table,
a list and a picture; the .docx also has auto-numbered headings, Symbol-font letters (µ, Ω) and a caption number.
  py make_samples.py      (needs pymupdf and python-docx)"""
import os, io
HERE = os.path.dirname(os.path.abspath(__file__))

def picture_png():
    import pymupdf as fitz
    d = fitz.open(); p = d.new_page(width=240, height=120)
    p.draw_rect(fitz.Rect(10, 10, 230, 110), color=(0.1, 0.3, 0.6), fill=(0.85, 0.9, 1), width=2)
    p.draw_circle((70, 60), 30, color=(0.8, 0.2, 0.2), fill=(1, 0.8, 0.8)); p.insert_text((120, 66), "Fig. A", fontsize=18)
    return p.get_pixmap(dpi=150).tobytes("png")

def make_pdf():
    import pymupdf as fitz
    doc = fitz.open(); png = picture_png(); total = 5
    for _ in range(total): doc.new_page()
    P = [doc[i] for i in range(total)]
    for i, p in enumerate(P):
        p.insert_text((40, 36), "SAMPLE MANUAL  |  Revision: B  |  Doc No. SM-100", fontsize=8)
        p.insert_text((40, 815), f"Example Corp. - Internal        Page {i + 1} of {total}", fontsize=8)
    P[0].insert_text((60, 300), "Sample Manual", fontsize=26)
    P[0].insert_text((60, 335), "Installation and test", fontsize=14)
    P[1].insert_text((60, 100), "Table of Contents", fontsize=16)
    for k, t in enumerate(["1 Scope ............................... 3", "2 Equipment ........................... 3",
                           "3 Procedure ........................... 4", "3.1 Mounting ........................ 4"]):
        P[1].insert_text((60, 140 + 22 * k), t, fontsize=11)
    y = 100
    def h(pg, t, size=15):
        nonlocal y; P[pg].insert_text((60, y), t, fontsize=size, fontname="hebo"); y += size * 2
    def para(pg, t):
        nonlocal y; r = fitz.Rect(60, y - 11, 540, y + 60); P[pg].insert_textbox(r, t, fontsize=10.5); y += 48
    h(2, "1 Scope"); para(2, "This manual describes how to install the sample unit and check that it works. It is written for technicians.")
    h(2, "2 Equipment"); para(2, "The following equipment is needed:")
    rows = [("Item", "Qty", "Note"), ("Screwdriver T10", "1", "torque 0.6 Nm"), ("Multimeter", "1", "calibrated"), ("Cable kit", "2", "")]
    for r_, row in enumerate(rows):
        for c, cell in enumerate(row):
            rect = fitz.Rect(60 + c * 150, y + r_ * 20, 210 + c * 150, y + 20 + r_ * 20)
            P[2].draw_rect(rect, width=0.6); P[2].insert_text((rect.x0 + 4, rect.y1 - 6), cell, fontsize=10, fontname="hebo" if r_ == 0 else "helv")
    y = 100
    h(3, "3 Procedure"); h(3, "3.1 Mounting", 12)
    for k, t in enumerate(["Switch the unit off.", "Fix the bracket with four screws.", "Connect the cable kit."]):
        P[3].insert_text((70, y), f"{k + 1}. {t}", fontsize=10.5); y += 18
    y += 8; P[3].insert_image(fitz.Rect(60, y, 300, y + 120), stream=png); y += 140
    P[3].insert_text((60, y), "Figure 1 - Bracket position", fontsize=9)
    doc.save(os.path.join(HERE, "sample.pdf")); print("written sample.pdf")

def make_docx():
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    d = Document()
    # auto-numbered headings: Heading 1/2 linked to a multilevel list (numbers are NOT in the text)
    numbering = d.part.numbering_part.element
    absn = OxmlElement("w:abstractNum"); absn.set(qn("w:abstractNumId"), "90")
    for lvl in range(2):
        l = OxmlElement("w:lvl"); l.set(qn("w:ilvl"), str(lvl))
        for tag, val in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1" if lvl == 0 else "%1.%2")):
            e = OxmlElement(tag); e.set(qn("w:val"), val); l.append(e)
        absn.append(l)
    numbering.insert(0, absn)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), "90"); a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), "90"); num.append(a); numbering.append(num)
    for i, sname in enumerate(("Heading 1", "Heading 2")):
        ppr = d.styles[sname].element.get_or_add_pPr(); npr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl"); il.set(qn("w:val"), str(i)); ni = OxmlElement("w:numId"); ni.set(qn("w:val"), "90")
        npr.append(il); npr.append(ni); ppr.insert(0, npr)
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "SAMPLE MANUAL  |  Revision: B"
    sec.footer.paragraphs[0].text = "Example Corp. - Internal"
    d.add_paragraph("Sample Manual", style="Title")
    # table of contents lines (style "toc 1"), skipped by default
    toc = d.styles.add_style("toc 1", 1)
    for t in ("Scope\t3", "Measurements\t3"): d.add_paragraph(t, style=toc)
    d.add_heading("Scope", 1)
    d.add_paragraph("This manual describes how to install the sample unit and check that it works.")
    d.add_heading("Measurements", 1)
    p = d.add_paragraph("The leakage current must be below 5 ")
    r = p.add_run("m"); r.font.name = "Symbol"; r.element.rPr.rFonts.set(qn("w:ascii"), "Symbol"); r.element.rPr.rFonts.set(qn("w:hAnsi"), "Symbol")
    p.add_run("A and the insulation above 10 M")
    r = p.add_run("W"); r.font.name = "Symbol"; r.element.rPr.rFonts.set(qn("w:ascii"), "Symbol"); r.element.rPr.rFonts.set(qn("w:hAnsi"), "Symbol")
    p.add_run(".")
    d.add_heading("Limits", 2)
    t = d.add_table(rows=3, cols=3); t.style = "Table Grid"
    for i, row in enumerate([("Test", "Min", "Max"), ("Voltage, V", "11.5", "12.5"), ("Current, mA", "-", "350")]):
        for j, v in enumerate(row): t.cell(i, j).text = v
    d.add_heading("Picture", 2)
    d.add_picture(io.BytesIO(picture_png()), width=Inches(2.4))
    cap = d.add_paragraph("Figure ", style="Caption")
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), " SEQ Figure \\* ARABIC ")
    rr = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = "1"; rr.append(tt); fld.append(rr); cap._p.append(fld)
    cap.add_run(" - Bracket position")
    d.add_paragraph("Check the bracket.", style="List Number"); d.add_paragraph("Tighten the screws.", style="List Number")
    d.save(os.path.join(HERE, "sample.docx")); print("written sample.docx")

if __name__ == "__main__":
    make_pdf(); make_docx()
